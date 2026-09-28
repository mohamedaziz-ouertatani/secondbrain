"""File → list of pages. A "page" is the unit a citation points to:
PDF page, PowerPoint slide, Word section (split at headings), or the whole .md/.txt note."""

import hashlib
import logging
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pymupdf
from docx.table import Table
from pptx import Presentation

from ..config import get_settings
from . import ocr

log = logging.getLogger(__name__)

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUPPORTED = {".pdf": PDF, ".pptx": PPTX, ".docx": DOCX, ".md": "text/markdown", ".txt": "text/plain"}

# Below this many non-space chars per page on average, treat the PDF as scanned.
MIN_CHARS_PER_PAGE = 30
# Pictures smaller than this (px) rarely hold readable text: icons, bullets, logos.
MIN_PICTURE = (300, 150)

# Metadata titles that say nothing about the content; fall back to the file name.
_GENERIC_TITLES = re.compile(
    r"^(powerpoint presentation|présentation powerpoint|presentation|présentation|diapositive \d+|slide \d+|"
    r"document|untitled|sans titre|microsoft word - .*|(template|modèle|modele)\b.*)$",
    re.IGNORECASE,
)


@dataclass
class Parsed:
    title: str
    mime: str
    pages: list[str]  # index 0 = page 1
    labels: list[str | None] | None = None  # optional per-page label, e.g. a Word section heading
    ocr_pages: set[int] = field(default_factory=set)  # 1-based pages/slides/sections holding OCR text

    @property
    def has_text(self) -> bool:
        chars = sum(len(re.sub(r"\s", "", p)) for p in self.pages)
        if self.mime != PDF:
            return chars > 0
        return chars >= MIN_CHARS_PER_PAGE * max(1, len(self.pages))


def normalize(text: str) -> str:
    # NFKC folds Arabic presentation forms (U+FB50–FEFF, common in PDF text) back to base letters,
    # and Latin ligatures like "ﬁ" to "fi". Without it, Arabic chunks embed and match poorly.
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("­", "")  # soft hyphen
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # word-\nwrap
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_title(title: str | None, path: Path) -> str:
    title = re.sub(r"^Microsoft (PowerPoint|Word) - ", "", (title or "").strip())
    title = re.sub(r"\.(pptx?|docx?)$", "", title, flags=re.IGNORECASE).strip()
    return path.stem if not title or _GENERIC_TITLES.match(title) else title


def parse(path: Path) -> Parsed:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"unsupported file type: {ext}")
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".docx":
        return _docx(path)
    text = normalize(path.read_text(encoding="utf-8", errors="replace"))
    title = path.stem
    if ext == ".md":
        m = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
        if m:
            title = m.group(1).strip()
    return Parsed(title=title, mime=SUPPORTED[ext], pages=[text])


def _worth_ocr(data: bytes, counts: Counter) -> bool:
    """Large enough to hold readable text, and not repeated in the file (template logos, banners)."""
    if counts[hashlib.sha1(data).hexdigest()] > 1:
        return False
    size = ocr.image_size(data)
    return size is not None and size[0] >= MIN_PICTURE[0] and size[1] >= MIN_PICTURE[1]


def _ocr_pictures(blobs: list[bytes], counts: Counter) -> str:
    return "\n".join(t for t in (ocr.ocr_image(b) for b in blobs if _worth_ocr(b, counts)) if t)


def _pdf(path: Path) -> Parsed:
    min_words = get_settings().ocr_min_words
    use_ocr = ocr.available()
    pages: list[str] = []
    ocr_pages: set[int] = set()
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if use_ocr and len(text.split()) < min_words and page.get_images():
                try:
                    extra = ocr.ocr_page(page)
                except Exception:  # keep the text layer if OCR fails
                    log.warning("OCR failed on %s page %d", path.name, i, exc_info=True)
                    extra = ""
                if extra:
                    text = f"{text}\n{extra}"
                    ocr_pages.add(i)
            pages.append(normalize(text))
        title = clean_title((doc.metadata or {}).get("title"), path)
    return Parsed(title=title, mime=PDF, pages=pages, ocr_pages=ocr_pages)


def _shape_texts(shapes) -> list[str]:
    """Text from shapes in reading order, recursing into groups; tables as 'a | b' rows."""
    out: list[str] = []
    for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            out.extend(_shape_texts(shape.shapes))
        elif shape.has_text_frame and shape.text_frame.text.strip():
            out.append(shape.text_frame.text)
        elif getattr(shape, "has_table", False) and shape.has_table:
            out.append("\n".join(" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows))
    return out


def _pictures(shapes) -> list[bytes]:
    out: list[bytes] = []
    for shape in shapes:
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            out.extend(_pictures(shape.shapes))
        elif shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            try:
                out.append(shape.image.blob)
            except Exception:  # noqa: BLE001, S112 -- linked or broken images have no blob
                continue
    return out


def _pptx(path: Path) -> Parsed:
    prs = Presentation(str(path))
    pages: list[str] = []
    use_ocr = ocr.available()
    slide_pics = [_pictures(slide.shapes) for slide in prs.slides] if use_ocr else []
    counts = Counter(hashlib.sha1(b).hexdigest() for pics in slide_pics for b in pics)
    ocr_pages: set[int] = set()
    for n, slide in enumerate(prs.slides, start=1):
        # Title first, then the rest top-to-bottom; paragraphs separated for the chunker.
        title_shape = slide.shapes.title
        parts = [title_shape.text_frame.text] if title_shape is not None and title_shape.has_text_frame else []
        # python-pptx returns a fresh proxy per access, so compare ids, not identity
        title_id = title_shape.shape_id if title_shape is not None else None
        parts += _shape_texts(s for s in slide.shapes if s.shape_id != title_id)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        if use_ocr:
            extra = _ocr_pictures(slide_pics[n - 1], counts)
            if extra:
                parts.append(extra)
                ocr_pages.add(n)
        pages.append(normalize("\n\n".join(p for p in parts if p.strip())))
    return Parsed(title=clean_title(prs.core_properties.title, path), mime=PPTX, pages=pages, ocr_pages=ocr_pages)


def _is_heading(style_name: str) -> bool:
    return bool(re.match(r"^(Heading [12]|Titre [12]|Title|Titre)$", style_name or ""))


def _para_images(block, part) -> list[bytes]:
    """Image bytes of the pictures drawn inside one paragraph."""
    out = []
    for blip in block._p.xpath(".//a:blip"):
        rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        if rid and rid in part.related_parts:
            out.append(part.related_parts[rid].blob)
    return out


def _docx(path: Path) -> Parsed:
    d = docx.Document(str(path))
    pages: list[str] = []
    labels: list[str | None] = []
    current: list[str] = []
    label: str | None = None
    first_heading: str | None = None
    use_ocr = ocr.available()
    blocks = list(d.iter_inner_content())
    # count per picture reference: a picture used twice points to one image part twice
    counts = Counter(hashlib.sha1(b).hexdigest() for blk in blocks if not isinstance(blk, Table)
                     for b in _para_images(blk, d.part)) if use_ocr else Counter()
    section_pics: list[bytes] = []
    ocr_pages: set[int] = set()

    def flush() -> None:
        extra = _ocr_pictures(section_pics, counts) if section_pics else ""
        section_pics.clear()
        if extra:
            current.append(extra)
        text = normalize("\n\n".join(current))
        if text:
            pages.append(text)
            labels.append(label)
            if extra:
                ocr_pages.add(len(pages))

    for block in blocks:
        if isinstance(block, Table):
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in block.rows]
            current.append("\n".join(rows))
            continue
        if use_ocr:
            section_pics.extend(_para_images(block, d.part))  # before the empty-text skip: picture paragraphs have no text
        text = block.text.strip()
        if not text:
            continue
        if _is_heading(block.style.name if block.style is not None else ""):
            flush()
            current, label = [text], text
            first_heading = first_heading or text
        else:
            current.append(text)
    flush()

    title = clean_title(d.core_properties.title, path)
    if title == path.stem and first_heading:
        title = first_heading
    return Parsed(title=title, mime=DOCX, pages=pages or [""], labels=labels or None, ocr_pages=ocr_pages)
