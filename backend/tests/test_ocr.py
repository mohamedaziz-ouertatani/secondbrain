import io

import pymupdf
import pytest

from app.ingest import ocr
from app.ingest.parse import parse

needs_data = pytest.mark.skipif(not ocr.available(), reason="Tesseract language data missing (python -m app.ingest.ocr --setup)")


def text_png(text: str, w: int = 400, h: int = 120) -> bytes:
    """A PNG of rendered text: an image with no text layer, like a screenshot."""
    doc = pymupdf.open()
    page = doc.new_page(width=w, height=h)
    page.insert_text((20, h / 2 + 10), text, fontsize=28)
    return page.get_pixmap(dpi=144).tobytes("png")


def test_clean_keeps_words_and_drops_junk():
    raw = "| 4”A 4 ny: As @\nEnglish French Translation Quality\n= |\ndocker compose up -d\nGNMT (RNN)\nBLEU 40"
    assert ocr.clean(raw) == "English French Translation Quality\ndocker compose up -d\nGNMT (RNN)"


@needs_data
def test_pdf_ocr_only_on_near_empty_image_pages(tmp_path):
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "This page has a real text layer about gaussian vectors " * 3, fontsize=10)
    p2 = doc.new_page()
    p2.insert_image(pymupdf.Rect(72, 72, 472, 192), stream=text_png("docker compose up"))
    f = tmp_path / "slides.pdf"
    doc.save(f)

    parsed = parse(f)
    assert parsed.ocr_pages == {2}
    assert "docker" in parsed.pages[1].lower() and "compose" in parsed.pages[1].lower()
    assert "gaussian" in parsed.pages[0] and "docker" not in parsed.pages[0]


@needs_data
def test_pptx_ocr_skips_small_and_repeated_pictures(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    unique, repeated = text_png("jenkins pipeline stage"), text_png("template footer banner")
    small = text_png("tiny logo", w=60, h=30)
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(io.BytesIO(unique), Inches(1), Inches(1))
    s1.shapes.add_picture(io.BytesIO(repeated), Inches(1), Inches(4))
    s1.shapes.add_picture(io.BytesIO(small), Inches(6), Inches(6))
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_picture(io.BytesIO(repeated), Inches(1), Inches(4))
    f = tmp_path / "deck.pptx"
    prs.save(f)

    parsed = parse(f)
    assert parsed.ocr_pages == {1}
    assert "jenkins" in parsed.pages[0].lower()
    assert "template" not in parsed.pages[0].lower() and "tiny" not in parsed.pages[0].lower()
    assert parsed.pages[1] == ""


@needs_data
def test_docx_picture_text_joins_its_section(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Workshop 4", level=1)
    d.add_paragraph("Run the following commands.")
    d.add_picture(io.BytesIO(text_png("docker build image")))
    f = tmp_path / "ws.docx"
    d.save(f)

    parsed = parse(f)
    assert parsed.ocr_pages == {1} and parsed.labels == ["Workshop 4"]
    assert "docker" in parsed.pages[0].lower()


def test_ocr_off_means_no_ocr_and_no_error(tmp_path, monkeypatch):
    from app import config

    monkeypatch.setenv("OCR_ENABLED", "false")
    config.get_settings.cache_clear()
    try:
        doc = pymupdf.open()
        doc.new_page().insert_image(pymupdf.Rect(72, 72, 472, 192), stream=text_png("docker compose up"))
        f = tmp_path / "img.pdf"
        doc.save(f)
        parsed = parse(f)
        assert parsed.ocr_pages == set() and parsed.pages == [""]
    finally:
        config.get_settings.cache_clear()
