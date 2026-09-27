from app.ingest.parse import normalize


def test_arabic_presentation_forms_are_folded():
    # "ﺍﻟﻔﻬﺮﺱ" in presentation forms, as often extracted from PDFs
    assert normalize("\ufe8d\ufedf\ufed4\ufeec\ufeae\ufeb1") == "الفهرس"


def test_hyphenation_and_ligatures():
    assert normalize("con-\nnexion ﬁable") == "connexion fiable"


# --- Office formats -------------------------------------------------------

from pathlib import Path

import docx
from pptx import Presentation
from pptx.util import Inches

from app.ingest.parse import DOCX, PPTX, clean_title, parse


def make_pptx(path: Path) -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Descente de gradient"
    s1.placeholders[1].text = "Optimization for ML, chapitre 2"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s2.shapes.title.text = "Pas d'apprentissage"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "Un pas trop grand fait diverger l'algorithme."
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "lr", "effet"
    table.cell(1, 0).text, table.cell(1, 1).text = "0.1", "diverge"
    s2.notes_slide.notes_text_frame.text = "Montrer la courbe de perte."
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide
    prs.core_properties.title = "Présentation PowerPoint"
    prs.save(path)


def test_pptx_one_page_per_slide(tmp_path):
    f = tmp_path / "cours-gradient.pptx"
    make_pptx(f)
    p = parse(f)
    assert p.mime == PPTX and p.has_text
    assert p.title == "cours-gradient"  # generic metadata title ignored
    assert len(p.pages) == 3  # blank slide kept so numbering matches the deck
    assert p.pages[0].startswith("Descente de gradient")
    assert all(page.count(t) == 1 for page, t in zip(p.pages, ["Descente de gradient", "Pas d'apprentissage"]))
    s2 = p.pages[1]
    assert s2.index("Pas d'apprentissage") < s2.index("Un pas trop grand") < s2.index("0.1 | diverge")
    assert s2.endswith("Notes: Montrer la courbe de perte.")
    assert p.pages[2] == ""


def test_docx_split_at_headings(tmp_path):
    d = docx.Document()
    d.add_paragraph("Intro before any heading.")
    d.add_heading("Loi normale", level=1)
    d.add_paragraph("La densité est symétrique autour de la moyenne.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "μ", "moyenne"
    d.add_heading("Loi de Poisson", level=2)
    d.add_paragraph("Modélise un nombre d'événements rares.")
    d.add_heading("Détail", level=3)  # level 3 stays inside the section
    d.add_paragraph("Espérance = variance = λ.")
    f = tmp_path / "proba.docx"
    d.save(f)

    p = parse(f)
    assert p.mime == DOCX
    assert p.title == "Loi normale"  # no metadata title → first heading
    assert p.labels == [None, "Loi normale", "Loi de Poisson"]
    assert p.pages[0] == "Intro before any heading."
    assert "μ | moyenne" in p.pages[1]
    assert "Détail" in p.pages[2] and "λ" in p.pages[2]


def test_clean_title():
    f = Path("chap1.pdf")
    assert clean_title("Microsoft PowerPoint - Chapitre 1.pptx", f) == "Chapitre 1"
    assert clean_title("Diapositive 1", f) == "chap1"
    assert clean_title("", f) == "chap1"
    assert clean_title("Big Data Analytics", f) == "Big Data Analytics"
