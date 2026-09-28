"""Integration: runs against a real Postgres (docker compose). Skipped if unreachable."""


import numpy as np


def _words(s: str) -> int:
    return len(s.split())


def fake_embed(texts):
    rng = [np.random.default_rng(abs(hash(t)) % 2**32) for t in texts]
    return [(v := r.standard_normal(1024).astype(np.float32)) / np.linalg.norm(v) for r in rng]


def chunk_count(db, rel):
    with db.get_pool().connection() as conn:
        return conn.execute(
            "SELECT count(c.id) AS n FROM documents d LEFT JOIN chunks c ON c.document_id = d.id WHERE d.path = %s", (rel,)
        ).fetchone()["n"]


def test_ingest_dedupe_replace_remove(env):
    from app.ingest.pipeline import ingest_file, rescan

    inbox, db = env
    course = inbox / "Algo"
    course.mkdir()
    note = course / "notes.md"
    note.write_text("# Sorting\n\n" + "\n\n".join(f"Paragraph {i} about quicksort pivots." for i in range(40)), encoding="utf-8")

    assert ingest_file(note, fake_embed, _words) == "ok"
    n1 = chunk_count(db, "Algo/notes.md")
    assert n1 > 0
    with db.get_pool().connection() as conn:
        doc = conn.execute("SELECT title, course FROM documents").fetchone()
    assert doc == {"title": "Sorting", "course": "Algo"}

    assert ingest_file(note, fake_embed, _words) == "skipped"
    assert chunk_count(db, "Algo/notes.md") == n1

    # a newer parser re-ingests unchanged files
    from app.ingest import pipeline
    pipeline.PARSER_VERSION += 1
    try:
        assert ingest_file(note, fake_embed, _words) == "ok"
        assert ingest_file(note, fake_embed, _words) == "skipped"
    finally:
        pipeline.PARSER_VERSION -= 1

    with db.get_pool().connection() as conn:
        first_seen = conn.execute("SELECT first_seen FROM documents").fetchone()["first_seen"]

    note.write_text("# Sorting\n\nShort now.", encoding="utf-8")
    assert ingest_file(note, fake_embed, _words) == "ok"
    assert chunk_count(db, "Algo/notes.md") == 1
    with db.get_pool().connection() as conn:  # re-ingesting never makes a file "new" again
        assert conn.execute("SELECT first_seen FROM documents").fetchone()["first_seen"] == first_seen

    note.unlink()
    assert rescan(fake_embed, _words)["removed"] == 1
    with db.get_pool().connection() as conn:
        assert conn.execute("SELECT count(*) AS n FROM chunks").fetchone()["n"] == 0


def test_embed_failure_records_error(env):
    from app.ingest.pipeline import ingest_file

    inbox, db = env
    f = inbox / "a.txt"
    f.write_text("some reasonably long text content for the test " * 3, encoding="utf-8")

    def boom(_):
        raise RuntimeError("ollama down")

    assert ingest_file(f, boom, _words) == "error"
    # an error is retried on next ingest even though the hash is unchanged
    assert ingest_file(f, fake_embed, _words) == "ok"


def test_docx_section_labels_stored(env, monkeypatch):
    import docx

    from app.ingest.pipeline import ingest_file
    from app.rag import retrieve

    inbox, db = env
    d = docx.Document()
    d.add_heading("Loi normale", level=1)
    d.add_paragraph("La densité est symétrique autour de la moyenne.")
    d.add_heading("Loi de Poisson", level=1)
    d.add_paragraph("Modélise un nombre d'événements rares.")
    (inbox / "Probability 2").mkdir()
    f = inbox / "Probability 2" / "lois.docx"
    d.save(f)

    assert ingest_file(f, fake_embed, _words) == "ok"
    monkeypatch.setattr(retrieve.ollama, "embed", fake_embed)
    hits = retrieve.dense("anything", k=10)
    assert sorted((h["page"], h["label"], h["course"]) for h in hits) == [
        (1, "Loi normale", "Probability 2"),
        (2, "Loi de Poisson", "Probability 2"),
    ]
