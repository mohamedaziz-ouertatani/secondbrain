import pymupdf
import pytest
from fakes import fake_embed, words


@pytest.fixture
def client(env):
    from fastapi.testclient import TestClient

    from app.main import create_app

    return TestClient(create_app())  # no `with`: no watcher


def test_pdf_page_images(env, client):
    from app.ingest.pipeline import rescan

    inbox, _ = env
    (inbox / "Optim").mkdir(parents=True)
    with pymupdf.open() as pdf:
        for n in (1, 2):
            pdf.new_page().insert_text((72, 72), f"Page {n}: the norm of u is the square root of the sum of squares. " * 3)
        pdf.save(inbox / "Optim" / "norms.pdf")
    (inbox / "Optim" / "note.md").write_text("A note about norms and the triangle inequality. " * 5, encoding="utf-8")
    rescan(fake_embed, words)

    ids = {d["path"]: d["id"] for d in client.get("/documents").json()}
    pdf_id = ids["Optim/norms.pdf"]

    page = client.get(f"/documents/{pdf_id}/pages/2.png")
    assert page.status_code == 200
    assert page.headers["content-type"] == "image/png"
    assert page.content.startswith(b"\x89PNG")

    assert client.get(f"/documents/{pdf_id}/pages/0.png").status_code == 404
    assert client.get(f"/documents/{pdf_id}/pages/3.png").status_code == 404
    assert client.get(f"/documents/{ids['Optim/note.md']}/pages/1.png").status_code == 404


def test_note_images_are_served_beside_the_note(env, client):
    from app.ingest.pipeline import rescan

    inbox, _ = env
    (inbox / "Optim" / "Serie 1").mkdir(parents=True)
    (inbox / "Optim" / "Serie 1.md").write_text("![fig](<Serie 1/fig 1.png>) " + "Norms and balls. " * 10, encoding="utf-8")
    (inbox / "Optim" / "Serie 1" / "fig 1.png").write_bytes(b"\x89PNG fake")
    (inbox / "Optim" / "Serie 1" / "notes.txt").write_text("not an image", encoding="utf-8")
    (inbox.parent / "outside.png").write_bytes(b"\x89PNG secret")
    rescan(fake_embed, words)
    doc = {d["path"]: d["id"] for d in client.get("/documents").json()}["Optim/Serie 1.md"]

    img = client.get(f"/documents/{doc}/assets/Serie%201/fig%201.png")
    assert img.status_code == 200 and img.headers["content-type"] == "image/png"
    assert img.content == b"\x89PNG fake"
    assert client.get(f"/documents/{doc}/assets/Serie%201/missing.png").status_code == 404
    assert client.get(f"/documents/{doc}/assets/Serie%201/notes.txt").status_code == 404  # images only
    assert client.get(f"/documents/{doc}/assets/..%2F..%2Foutside.png").status_code == 404  # stays in the inbox


def _deck(path, n):
    from pptx import Presentation

    prs = Presentation()
    for i in range(1, n + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i}"
        slide.placeholders[1].text = f"Gradient descent step {i} moves against the gradient. " * 3
    prs.save(path)


@pytest.fixture
def fake_soffice(monkeypatch, tmp_path_factory):
    """Stand-in for LibreOffice: one PDF page per slide, and a count of conversions."""
    from pptx import Presentation

    from app.ingest import slides

    calls = []

    def convert(src, out):
        calls.append(src)
        with pymupdf.open() as pdf:
            for _ in Presentation(str(src)).slides:
                pdf.new_page(width=720, height=405)
            pdf.save(out)

    cache = tmp_path_factory.mktemp("slides")
    monkeypatch.setattr(slides, "cache_dir", lambda: cache)
    monkeypatch.setattr(slides, "converter", lambda: "soffice")
    monkeypatch.setattr(slides, "convert", convert)
    return calls


def test_pptx_slide_images(env, client, fake_soffice):
    from app.ingest.pipeline import rescan

    inbox, _ = env
    (inbox / "ML").mkdir(parents=True)
    _deck(inbox / "ML" / "descent.pptx", 3)
    rescan(fake_embed, words)
    doc = {d["path"]: d["id"] for d in client.get("/documents").json()}["ML/descent.pptx"]

    assert client.get(f"/documents/{doc}/pages").json()["page_images"] is True
    for n in (1, 3):
        img = client.get(f"/documents/{doc}/pages/{n}.png")
        assert img.status_code == 200 and img.content.startswith(b"\x89PNG")
    assert client.get(f"/documents/{doc}/pages/4.png").status_code == 404
    assert len(fake_soffice) == 1  # converted once, then served from the cache


def test_slides_are_text_only_without_libreoffice(env, client, monkeypatch):
    from app.ingest import slides
    from app.ingest.pipeline import rescan

    monkeypatch.setattr(slides, "converter", lambda: None)
    inbox, _ = env
    (inbox / "ML").mkdir(parents=True)
    _deck(inbox / "ML" / "descent.pptx", 2)
    (inbox / "ML" / "note.md").write_text("A note about gradients and step sizes. " * 5, encoding="utf-8")
    rescan(fake_embed, words)
    ids = {d["path"]: d["id"] for d in client.get("/documents").json()}

    assert client.get(f"/documents/{ids['ML/descent.pptx']}/pages").json()["page_images"] is False
    assert client.get(f"/documents/{ids['ML/descent.pptx']}/pages/1.png").status_code == 404
    assert client.get(f"/documents/{ids['ML/note.md']}/pages").json()["page_images"] is False
